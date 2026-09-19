"""
GSTR-2B Consolidation Engine
-----------------------------
Implements the exact output format of the reference workbook:
  - Per-type sheets (B2B, B2BA, B2B-CDNR, IMPG, IMPGSEZ, IMPGA, + any other
    GSTR-2B sheet type present) with a "Month" column prepended.
  - RCM: reverse-charge rows pulled out of B2B + B2BA (and removed from
    those sheets), unioned into one sheet with a "Type" column.
  - Net: B2B + CDNR + IMPG + IMPGSEZ + B2BA + IMPGA mapped into one common
    schema (B2B's own column set), for a single "net ITC register".
  - Pivot: Net grouped by Month, IGST/CGST/SGST/Cess totals, Jan-Dec order
    + Grand Total row.
"""

import io
from openpyxl.utils import get_column_letter
import openpyxl

READ_ME_SHEET = "Read me"

MONTH_ORDER = ["January", "February", "March", "April", "May", "June",
               "July", "August", "September", "October", "November", "December"]

# The common "Net" schema — this is B2B's own header list, unprefixed by "Month"
# (Month is added separately as the 2nd column, "Sheet" as the 1st).
NET_SCHEMA = [
    "GSTIN of supplier", "Trade/Legal name", "Invoice number", "Invoice type",
    "Invoice Date", "Invoice Value(₹)", "Place of supply",
    "Supply Attract Reverse Charge", "Taxable Value (₹)", "Integrated Tax(₹)",
    "Central Tax(₹)", "State/UT Tax(₹)", "Cess(₹)",
    "GSTR-1/1A/IFF/GSTR-5 Period", "GSTR-1/1A/IFF/GSTR-5 Filing Date",
    "ITC Availability", "Reason", "Applicable % of Tax Rate", "Source",
    "IRN", "IRN Date", "GSTR-1/IFF/GSTR-5 Period", "GSTR-1/IFF/GSTR-5 Filing Date",
]


# --------------------------------------------------------------------------
# Read me sheet parsing
# --------------------------------------------------------------------------

def get_period_info(wb):
    """Read Tax Period (used as 'Month') + Financial Year from 'Read me'."""
    if READ_ME_SHEET not in wb.sheetnames:
        return None, None
    ws = wb[READ_ME_SHEET]
    fy, period = None, None
    for row in ws.iter_rows(min_row=1, max_row=20):
        for cell in row:
            if cell.value and isinstance(cell.value, str):
                label = cell.value.strip().lower()
                if label == "financial year":
                    fy = ws.cell(row=cell.row, column=3).value
                elif label == "tax period":
                    period = ws.cell(row=cell.row, column=3).value
    return period, fy


def get_extra_info(wb):
    info = {}
    if READ_ME_SHEET not in wb.sheetnames:
        return info
    ws = wb[READ_ME_SHEET]
    wanted = {"gstin", "legal name", "trade name (if any)", "date of generation"}
    for row in ws.iter_rows(min_row=1, max_row=20):
        for cell in row:
            if cell.value and isinstance(cell.value, str):
                label = cell.value.strip().lower()
                if label in wanted:
                    info[cell.value.strip()] = ws.cell(row=cell.row, column=3).value
    return info


# --------------------------------------------------------------------------
# Generic sheet extraction (handles the variable-depth merged headers)
# --------------------------------------------------------------------------

def detect_header_end_row(ws):
    max_col = ws.max_column
    candidates = [
        m.max_row
        for m in ws.merged_cells.ranges
        if 5 <= m.min_row <= 7 and not (m.min_col == 1 and m.max_col == max_col)
    ]
    return max(candidates) if candidates else 6


def build_merge_map(ws, max_row):
    m = {}
    max_col = ws.max_column
    for rng in ws.merged_cells.ranges:
        if rng.min_row > max_row:
            continue
        if rng.min_col == 1 and rng.max_col == max_col:
            continue
        anchor_val = ws.cell(row=rng.min_row, column=rng.min_col).value
        for r in range(rng.min_row, min(rng.max_row, max_row) + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                m[(r, c)] = anchor_val
    return m


def get_cell(ws, merge_map, row, col):
    if (row, col) in merge_map:
        return merge_map[(row, col)]
    return ws.cell(row=row, column=col).value


def dedupe_headers(headers):
    seen = {}
    out = []
    for h in headers:
        if h not in seen:
            seen[h] = 1
            out.append(h)
        else:
            seen[h] += 1
            out.append(f"{h}_{seen[h] - 1}")
    return out


def extract_sheet(ws):
    """Return (headers, data_rows) for a single GSTR-2B data sheet."""
    header_end = detect_header_end_row(ws)
    merge_map = build_merge_map(ws, header_end)
    max_col = ws.max_column

    headers = []
    for c in range(1, max_col + 1):
        val = None
        for r in range(header_end, 4, -1):
            v = get_cell(ws, merge_map, r, c)
            if v not in (None, ""):
                val = str(v).strip()
                break
        headers.append(val or f"Column_{c}")
    headers = dedupe_headers(headers)

    data_rows = []
    for r in range(header_end + 1, ws.max_row + 1):
        row_vals = [ws.cell(row=r, column=c).value for c in range(1, max_col + 1)]
        if any(v not in (None, "") for v in row_vals):
            data_rows.append(row_vals)

    return headers, data_rows


# --------------------------------------------------------------------------
# Per-type consolidation (raw, before RCM split-out)
# --------------------------------------------------------------------------

def collect_raw_sheets(uploaded_files):
    """
    Returns:
      raw: { sheet_name: {"headers": ["Month", ...], "rows": [[...]]} }
      file_summaries: [...]
      warnings: [...]
    """
    raw = {}
    file_summaries = []
    warnings = []

    for uf in uploaded_files:
        try:
            wb = openpyxl.load_workbook(io.BytesIO(uf.getvalue()), data_only=True)
        except Exception as e:
            warnings.append(f"Could not open '{uf.name}': {e}")
            continue

        month, fy = get_period_info(wb)
        extra = get_extra_info(wb)
        sheet_names = [s for s in wb.sheetnames if s != READ_ME_SHEET]
        rows_added = 0

        for sheet_name in sheet_names:
            ws = wb[sheet_name]
            try:
                headers, data_rows = extract_sheet(ws)
            except Exception as e:
                warnings.append(f"Skipped sheet '{sheet_name}' in '{uf.name}': {e}")
                continue

            if sheet_name not in raw:
                raw[sheet_name] = {"headers": headers, "rows": []}

            ref_headers = raw[sheet_name]["headers"]

            if headers == ref_headers:
                aligned_rows = [list(r) for r in data_rows]
            else:
                name_to_idx = {}
                for idx, h in enumerate(headers):
                    name_to_idx.setdefault(h, idx)
                aligned_rows = []
                for row in data_rows:
                    new_row = [
                        row[name_to_idx[h]] if h in name_to_idx and name_to_idx[h] < len(row) else None
                        for h in ref_headers
                    ]
                    aligned_rows.append(new_row)
                extra_cols = [h for h in headers if h not in ref_headers]
                if extra_cols:
                    ref_headers = ref_headers + extra_cols
                    raw[sheet_name]["headers"] = ref_headers
                    for r in raw[sheet_name]["rows"]:
                        r.extend([None] * len(extra_cols))
                    for i, row in enumerate(data_rows):
                        for h in extra_cols:
                            idx = name_to_idx.get(h)
                            aligned_rows[i].append(row[idx] if idx is not None and idx < len(row) else None)

            for row in aligned_rows:
                raw[sheet_name]["rows"].append([month] + list(row))
                rows_added += 1

        file_summaries.append({
            "filename": uf.name,
            "period": month,
            "financial_year": fy,
            "gstin": extra.get("GSTIN", ""),
            "legal_name": extra.get("Legal Name", ""),
            "sheets": len(sheet_names),
            "rows": rows_added,
        })

    for payload in raw.values():
        payload["headers"] = ["Month"] + payload["headers"]

    return raw, file_summaries, warnings


# --------------------------------------------------------------------------
# RCM split-out (B2B + B2BA reverse-charge rows)
# --------------------------------------------------------------------------

def split_rcm(raw):
    """
    Mutates raw in place: removes reverse-charge rows from B2B and B2BA.
    Returns the RCM sheet payload {"headers": [...], "rows": [...]}.
    """
    rcm_rows = []
    b2b_headers = raw.get("B2B", {}).get("headers", [])
    b2ba_headers = raw.get("B2BA", {}).get("headers", [])

    # RCM schema: Type, then B2B's own headers (minus B2B's own "Month" dup),
    # then B2BA's amendment-specific columns that carry genuinely distinct
    # information (the revised invoice #/date, the ITC-reduction flag, the
    # delta tax amounts, and Remarks). B2BA also repeats near-duplicate
    # "GSTR period / filing date" columns (with a differently-ordered name
    # like "GSTR-1/IFF/1A/..." vs B2B's "GSTR-1/1A/IFF/..."); these are the
    # same concept as B2B's own period columns so they're deliberately left
    # out here rather than added as extra duplicate columns.
    b2b_cols = [h for h in b2b_headers if h != "Month"]
    b2ba_cols = [h for h in b2ba_headers if h != "Month"]
    rcm_extra_b2ba_cols = [
        "Invoice number_1", "Invoice Date_1",
        "Whether ITC to be reduced (Taxpayer's Input)",
        "Integrated Tax(₹)_1", "Central Tax(₹)_1", "State/UT Tax(₹)_1", "Cess(₹)_1",
        "Remarks",
    ]
    extra_b2ba_cols = [h for h in rcm_extra_b2ba_cols if h in b2ba_cols]
    rcm_headers = ["Type", "Month"] + b2b_cols + extra_b2ba_cols

    def move_rows(sheet_name, type_label, own_cols):
        payload = raw.get(sheet_name)
        if not payload:
            return
        headers = payload["headers"]
        try:
            rc_idx = headers.index("Supply Attract Reverse Charge")
        except ValueError:
            return
        keep_rows = []
        for row in payload["rows"]:
            if str(row[rc_idx]).strip().lower() == "yes":
                month_val = row[0]
                col_map = dict(zip(headers[1:], row[1:]))  # skip Month
                new_row = ["Type_PLACEHOLDER", month_val]
                for h in b2b_cols:
                    new_row.append(col_map.get(h))
                for h in extra_b2ba_cols:
                    new_row.append(col_map.get(h))
                new_row[0] = type_label
                rcm_rows.append(new_row)
            else:
                keep_rows.append(row)
        payload["rows"] = keep_rows

    move_rows("B2B", "B2B", b2b_cols)
    move_rows("B2BA", "B2BA", b2ba_cols)

    return {"headers": rcm_headers, "rows": rcm_rows}


# --------------------------------------------------------------------------
# Net sheet (common-schema union of the core ITC-bearing sheets)
# --------------------------------------------------------------------------

def _colmap(headers, row):
    return dict(zip(headers[1:], row[1:]))  # skip Month at index 0


def _to_num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def build_net_rows(raw):
    net_rows = []

    # --- B2B: direct 1:1 mapping ---
    if "B2B" in raw:
        headers = raw["B2B"]["headers"]
        for row in raw["B2B"]["rows"]:
            m = _colmap(headers, row)
            net_rows.append(["B2B", row[0]] + [m.get(c) for c in NET_SCHEMA])

    # --- B2BA: use the revised ("_1") invoice number/date + the main (non-_1) tax
    #     amounts; GSTR filing / ITC-availability columns are left blank, matching
    #     the reference workbook (amendment rows only carry core invoice fields).
    if "B2BA" in raw:
        headers = raw["B2BA"]["headers"]
        for row in raw["B2BA"]["rows"]:
            m = _colmap(headers, row)
            mapped = {
                "GSTIN of supplier": m.get("GSTIN of supplier"),
                "Trade/Legal name": m.get("Trade/Legal name"),
                "Invoice number": m.get("Invoice number_1", m.get("Invoice number")),
                "Invoice type": m.get("Invoice type"),
                "Invoice Date": m.get("Invoice Date_1", m.get("Invoice Date")),
                "Invoice Value(₹)": m.get("Invoice Value(₹)"),
                "Place of supply": m.get("Place of supply"),
                "Supply Attract Reverse Charge": m.get("Supply Attract Reverse Charge"),
                "Taxable Value (₹)": m.get("Taxable Value (₹)"),
                "Integrated Tax(₹)": m.get("Integrated Tax(₹)"),
                "Central Tax(₹)": m.get("Central Tax(₹)"),
                "State/UT Tax(₹)": m.get("State/UT Tax(₹)"),
                "Cess(₹)": m.get("Cess(₹)"),
            }
            net_rows.append(["B2BA", row[0]] + [mapped.get(c) for c in NET_SCHEMA])

    # --- CDNR (B2B-CDNR): Note fields -> Invoice fields; Credit Notes get their
    #     taxable value / tax amounts NEGATED (they reduce ITC), Debit Notes stay
    #     positive (they add to ITC).
    if "B2B-CDNR" in raw:
        headers = raw["B2B-CDNR"]["headers"]
        for row in raw["B2B-CDNR"]["rows"]:
            m = _colmap(headers, row)
            note_type = (m.get("Note type") or "").strip().lower()
            sign = -1 if note_type == "credit note" else 1
            mapped = {
                "GSTIN of supplier": m.get("GSTIN of supplier"),
                "Trade/Legal name": m.get("Trade/Legal name"),
                "Invoice number": m.get("Note number"),
                "Invoice type": m.get("Note type"),
                "Invoice Date": m.get("Note date"),
                "Invoice Value(₹)": m.get("Note Value (₹)") or m.get("Note Value  (₹)"),
                "Place of supply": m.get("Place of supply"),
                "Supply Attract Reverse Charge": m.get("Supply Attract Reverse Charge"),
                "Taxable Value (₹)": sign * _to_num(m.get("Taxable Value (₹)")),
                "Integrated Tax(₹)": sign * _to_num(m.get("Integrated Tax(₹)")),
                "Central Tax(₹)": sign * _to_num(m.get("Central Tax(₹)")),
                "State/UT Tax(₹)": sign * _to_num(m.get("State/UT Tax(₹)")),
                "Cess(₹)": sign * _to_num(m.get("Cess(₹)")),
            }
            net_rows.append(["CDNR", row[0]] + [mapped.get(c) for c in NET_SCHEMA])

    # --- IMPG / IMPGA: no real counterparty GSTIN -> constants "IMPG" /
    #     "IMPORT PARTY"; Invoice number = Port Code; Invoice Date = Bill of
    #     Entry Date; Invoice Value = Taxable Value + Integrated Tax + Cess;
    #     Central/State Tax and Cess are not carried into Net (imports only
    #     attract IGST for this mapping).
    for sheet_name, type_label in (("IMPG", "IMPG"), ("IMPGA", "IMPGA")):
        if sheet_name not in raw:
            continue
        headers = raw[sheet_name]["headers"]
        for row in raw[sheet_name]["rows"]:
            m = _colmap(headers, row)
            taxable = _to_num(m.get("Taxable Value"))
            igst = _to_num(m.get("Integrated Tax(₹)"))
            cess = _to_num(m.get("Cess(₹)"))
            mapped = {
                "GSTIN of supplier": "IMPG",
                "Trade/Legal name": "IMPORT PARTY",
                "Invoice number": m.get("Port Code"),
                "Invoice type": "IMPG",
                "Invoice Date": m.get("Date"),
                "Invoice Value(₹)": taxable + igst + cess,
                "Taxable Value (₹)": taxable,
                "Integrated Tax(₹)": igst,
            }
            net_rows.append([type_label, row[0]] + [mapped.get(c) for c in NET_SCHEMA])

    # --- IMPGSEZ: has a real supplier GSTIN/name; same Invoice number/date/value
    #     logic as IMPG.
    if "IMPGSEZ" in raw:
        headers = raw["IMPGSEZ"]["headers"]
        for row in raw["IMPGSEZ"]["rows"]:
            m = _colmap(headers, row)
            taxable = _to_num(m.get("Taxable Value"))
            igst = _to_num(m.get("Integrated Tax(₹)"))
            cess = _to_num(m.get("Cess(₹)"))
            mapped = {
                "GSTIN of supplier": m.get("GSTIN of supplier"),
                "Trade/Legal name": m.get("Trade/Legal name"),
                "Invoice number": m.get("Port Code"),
                "Invoice type": "IMPGSEZ",
                "Invoice Date": m.get("Date"),
                "Invoice Value(₹)": taxable + igst + cess,
                "Taxable Value (₹)": taxable,
                "Integrated Tax(₹)": igst,
            }
            net_rows.append(["IMPGSEZ", row[0]] + [mapped.get(c) for c in NET_SCHEMA])

    return ["Sheet", "Month"] + NET_SCHEMA, net_rows


# --------------------------------------------------------------------------
# Pivot (Net grouped by Month)
# --------------------------------------------------------------------------

def build_pivot(net_headers, net_rows):
    idx_month = net_headers.index("Month")
    idx_igst = net_headers.index("Integrated Tax(₹)")
    idx_cgst = net_headers.index("Central Tax(₹)")
    idx_sgst = net_headers.index("State/UT Tax(₹)")
    idx_cess = net_headers.index("Cess(₹)")

    sums = {}
    for row in net_rows:
        month = row[idx_month]
        if not month:
            continue
        if month not in sums:
            sums[month] = [0.0, 0.0, 0.0, 0.0]
        sums[month][0] += _to_num(row[idx_igst])
        sums[month][1] += _to_num(row[idx_cgst])
        sums[month][2] += _to_num(row[idx_sgst])
        sums[month][3] += _to_num(row[idx_cess])

    headers = ["Row Labels", "Sum of Integrated Tax(₹)", "Sum of Central Tax(₹)",
               "Sum of State/UT Tax(₹)", "Sum of Cess(₹)"]
    rows = []
    ordered_months = [m for m in MONTH_ORDER if m in sums]
    # append any month names that don't match the standard 12 (unexpected data) at the end
    ordered_months += [m for m in sums if m not in MONTH_ORDER]

    grand = [0.0, 0.0, 0.0, 0.0]
    for month in ordered_months:
        vals = sums[month]
        rows.append([month] + vals)
        for i in range(4):
            grand[i] += vals[i]
    rows.append(["Grand Total"] + grand)

    return headers, rows


# --------------------------------------------------------------------------
# Top-level orchestration
# --------------------------------------------------------------------------

def process_files(uploaded_files):
    raw, file_summaries, warnings = collect_raw_sheets(uploaded_files)

    rcm_payload = split_rcm(raw)
    net_headers, net_rows = build_net_rows(raw)
    pivot_headers, pivot_rows = build_pivot(net_headers, net_rows)

    result_sheets = {}

    # Ordered "primary" sheets first, matching the reference workbook order.
    primary_order = ["RCM", "B2B", "B2BA", "B2B-CDNR", "Pivot", "Net",
                      "IMPG", "IMPGSEZ", "IMPGA"]

    result_sheets["RCM"] = rcm_payload
    for name in ["B2B", "B2BA", "B2B-CDNR"]:
        if name in raw:
            result_sheets[name] = raw[name]
    result_sheets["Pivot"] = {"headers": pivot_headers, "rows": pivot_rows}
    result_sheets["Net"] = {"headers": net_headers, "rows": net_rows}
    for name in ["IMPG", "IMPGSEZ", "IMPGA"]:
        if name in raw:
            result_sheets[name] = raw[name]

    # Any other GSTR-2B sheet types present (ITC Available, ISD, ECO, rejected
    # / amendment variants, etc.) are appended after, unchanged, so nothing
    # from the uploaded files is lost.
    for name, payload in raw.items():
        if name not in result_sheets:
            result_sheets[name] = payload

    ordered_result = {}
    for name in primary_order:
        if name in result_sheets:
            ordered_result[name] = result_sheets[name]
    for name, payload in result_sheets.items():
        if name not in ordered_result:
            ordered_result[name] = payload

    return ordered_result, file_summaries, warnings
