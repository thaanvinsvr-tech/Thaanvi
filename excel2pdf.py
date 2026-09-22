"""
Excel -> PDF Converter
------------------------
Renders every sheet of an uploaded workbook as a table in a downloadable
PDF (landscape, one section per sheet, auto-paginated for long sheets).
"""

import io

import openpyxl
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle


def _fmt(v):
    if v is None:
        return ""
    return str(v)


def _table_font_size(n_cols):
    if n_cols <= 6:
        return 9
    if n_cols <= 10:
        return 7.5
    if n_cols <= 16:
        return 6
    return 5


def workbook_to_pdf(file_bytes, filename):
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)

    buf = io.BytesIO()
    page_size = landscape(A4)
    doc = SimpleDocTemplate(
        buf, pagesize=page_size,
        leftMargin=10 * mm, rightMargin=10 * mm, topMargin=12 * mm, bottomMargin=12 * mm,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("SheetTitle", parent=styles["Heading2"], spaceAfter=6)

    usable_width = page_size[0] - 20 * mm

    story = []
    first = True
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        rows = [list(r) for r in rows if any(v not in (None, "") for v in r)]
        if not rows:
            continue

        if not first:
            story.append(PageBreak())
        first = False

        story.append(Paragraph(f"{filename} — {sheet_name}", title_style))
        story.append(Spacer(1, 4))

        n_cols = max(len(r) for r in rows)
        font_size = _table_font_size(n_cols)
        col_width = usable_width / n_cols

        data = [[_fmt(v) for v in (r + [None] * (n_cols - len(r)))] for r in rows]

        # cap total rows-per-table isn't needed; reportlab Table auto-splits
        # across pages via the doc's story when it doesn't fit.
        tbl = Table(data, colWidths=[col_width] * n_cols, repeatRows=1)
        tbl.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), font_size),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F2F2")]),
        ]))
        story.append(tbl)

    if not story:
        story = [Paragraph("No data found in this workbook.", styles["Normal"])]

    doc.build(story)
    buf.seek(0)
    return buf
