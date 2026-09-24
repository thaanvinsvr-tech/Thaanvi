"""
Thaanvi's Audit Suite
------------------------
A GST + Audit toolkit, organised into categories (GST Compliance, QRMP,
Reconciliation, File Conversion, Audit Tools). Only tools that actually
have working logic behind them are clickable — everything else is shown
as "Coming Soon" so the navigation structure is ready for future modules
without pretending unbuilt tools work.

This file is presentation only. All actual GST/file processing lives in
engine.py, engine3b.py, pdf2excel.py, excel2pdf.py and json2excel.py, and
none of that logic is touched here.
"""

import io
from datetime import datetime

import streamlit as st
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

import engine as engine2b
import engine3b
import pdf2excel
import excel2pdf
import json2excel

st.set_page_config(page_title="Thaanvi's Audit Suite", page_icon="\U0001F9FE", layout="wide")


# ==========================================================================
# Theme — navy sidebar (fixed brand identity) + togglable light/dark content
# ==========================================================================

NAVY = "#17324D"
NAVY_SOFT = "#22415F"
TEAL = "#0F9D8A"

CONTENT_LIGHT = {
    "bg": "#F6F8FA", "card_bg": "#FFFFFF", "text": "#1F2937",
    "muted": "#5B6B7C", "border": "#E2E8F0",
}
CONTENT_DARK = {
    "bg": "#0E1620", "card_bg": "#16212C", "text": "#E5E9EE",
    "muted": "#9FB0C0", "border": "#2A3846",
}


def inject_theme(mode):
    c = CONTENT_DARK if mode == "Dark" else CONTENT_LIGHT
    st.markdown(f"""
    <style>
    .stApp {{ background-color: {c['bg']}; }}
    .stApp, .stApp p, .stApp span, .stApp div, .stApp label {{ color: {c['text']}; }}
    h1, h2, h3, h4, h5, h6 {{ color: {c['text']} !important; }}

    section[data-testid="stSidebar"] {{ background-color: {NAVY}; }}
    section[data-testid="stSidebar"] * {{ color: #FFFFFF !important; }}
    section[data-testid="stSidebar"] .stCaption, section[data-testid="stSidebar"] small {{
        color: #C7D5E3 !important;
    }}
    section[data-testid="stSidebar"] hr {{ border-color: #33506E; }}

    section[data-testid="stSidebar"] .stButton>button {{
        background-color: transparent; border: 1px solid transparent;
        text-align: left; justify-content: flex-start; border-radius: 8px;
        font-weight: 500; padding: 0.45rem 0.7rem;
    }}
    section[data-testid="stSidebar"] .stButton>button:hover {{
        background-color: {NAVY_SOFT}; border-color: {NAVY_SOFT};
    }}
    section[data-testid="stSidebar"] .stButton>button[kind="primary"] {{
        background-color: {TEAL}; border-color: {TEAL}; color: #FFFFFF !important;
    }}
    section[data-testid="stSidebar"] .stButton>button[kind="primary"] * {{ color: #FFFFFF !important; }}
    section[data-testid="stSidebar"] .stButton>button:disabled {{
        background-color: transparent; color: #6E859B !important; opacity: 0.6;
    }}

    div[data-testid="stAppViewContainer"] .stButton>button,
    div[data-testid="stAppViewContainer"] .stDownloadButton>button {{
        background-color: {TEAL}; color: #FFFFFF; border-radius: 8px;
        border: 1px solid {TEAL}; font-weight: 600;
    }}
    div[data-testid="stAppViewContainer"] .stButton>button:hover,
    div[data-testid="stAppViewContainer"] .stDownloadButton>button:hover {{ opacity: 0.9; }}
    div[data-testid="stAppViewContainer"] .stButton>button[kind="secondary"] {{
        background-color: {c['card_bg']}; color: {NAVY}; border: 1px solid {c['border']};
    }}
    div[data-testid="stAppViewContainer"] .stButton>button:disabled {{
        background-color: {c['bg']}; color: {c['muted']} !important;
        border: 1px dashed {c['border']}; opacity: 0.85;
    }}

    div[data-testid="stMetric"] {{
        background-color: {c['card_bg']}; padding: 16px; border-radius: 12px;
        border: 1px solid {c['border']}; border-top: 3px solid {TEAL};
    }}
    div[data-testid="stMetric"] label {{ color: {c['muted']} !important; }}
    div[data-testid="stVerticalBlockBorderWrapper"] {{
        background-color: {c['card_bg']}; border-radius: 12px; border-color: {c['border']} !important;
    }}

    div[data-testid="stFileUploaderDropzone"] {{
        background-color: {c['card_bg']}; border: 1.5px dashed {TEAL}; border-radius: 12px;
    }}

    .stTabs [data-baseweb="tab"] {{ color: {c['muted']}; font-weight: 500; }}
    .stTabs [aria-selected="true"] {{ color: {TEAL} !important; border-bottom-color: {TEAL} !important; }}

    div[data-testid="stAlert"] {{ border-radius: 10px; }}

    .gcs-topbar {{ display: flex; align-items: baseline; gap: 10px; margin-bottom: 0.25rem; }}
    .gcs-topbar-title {{ font-size: 1.05rem; font-weight: 700; color: {NAVY if mode != "Dark" else "#FFFFFF"}; }}
    .gcs-topbar-tag {{ font-size: 0.85rem; color: {c['muted']}; }}

    .gcs-hero-title {{ font-size: 1.9rem; font-weight: 800; margin-bottom: 0; }}
    .gcs-hero-tagline {{ font-size: 1.05rem; color: {TEAL}; font-weight: 600; margin: 4px 0 6px 0; }}
    .gcs-hero-desc {{ font-size: 0.95rem; color: {c['muted']}; margin-bottom: 1rem; }}

    .gcs-soon-badge {{
        display: inline-block; font-size: 0.7rem; font-weight: 700; letter-spacing: 0.5px;
        background-color: {c['border']}; color: {c['muted']}; padding: 2px 8px;
        border-radius: 999px; margin-left: 6px;
    }}
    </style>
    """, unsafe_allow_html=True)


theme_mode = st.session_state.get("theme_mode_radio", "Light")
inject_theme(theme_mode)


def top_bar():
    st.markdown(
        '<div class="gcs-topbar">'
        '<span class="gcs-topbar-title">\U0001F9FE Thaanvi\'s Audit Suite</span>'
        '<span class="gcs-topbar-tag">Simplify. Reconcile. Review. Report.</span>'
        '</div><hr style="margin-top:6px;">',
        unsafe_allow_html=True,
    )


def tool_hero(icon, title, tagline, description):
    st.markdown(f'<div class="gcs-hero-title">{icon} {title}</div>', unsafe_allow_html=True)
    st.markdown(f'<div class="gcs-hero-tagline">{tagline}</div>', unsafe_allow_html=True)
    st.markdown(f'<div class="gcs-hero-desc">{description}</div>', unsafe_allow_html=True)


# ==========================================================================
# Shared helpers (output-file styling — unrelated to on-screen UI, unchanged)
# ==========================================================================

HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(bold=True, size=14, color="1F4E78")
SUBTITLE_FONT = Font(italic=True, size=10, color="666666")
INVALID_SHEET_CHARS = str.maketrans({c: "-" for c in r'\/?*[]:'})


def safe_sheet_name(name):
    return name.translate(INVALID_SHEET_CHARS)[:31]


# ==========================================================================
# GSTR-2B tool  (logic calls unchanged from the previous version)
# ==========================================================================

HEADER_FILL_2B = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")


def build_2b_workbook(result_sheets, file_summaries):
    out_wb = openpyxl.Workbook()
    out_wb.remove(out_wb.active)

    ws = out_wb.create_sheet("Consolidation Summary")
    ws["A1"] = "GSTR-2B Consolidation Summary"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Generated on {datetime.now().strftime('%d-%b-%Y %H:%M')}"
    ws["A2"].font = SUBTITLE_FONT

    headers = ["File Name", "Month", "Financial Year", "GSTIN", "Legal Name", "Sheets Found", "Rows Consolidated"]
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=c, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL_2B
    for r, fs in enumerate(file_summaries, start=5):
        ws.cell(row=r, column=1, value=fs["filename"])
        ws.cell(row=r, column=2, value=fs["period"])
        ws.cell(row=r, column=3, value=fs["financial_year"])
        ws.cell(row=r, column=4, value=fs["gstin"])
        ws.cell(row=r, column=5, value=fs["legal_name"])
        ws.cell(row=r, column=6, value=fs["sheets"])
        ws.cell(row=r, column=7, value=fs["rows"])
    for c, width in zip(range(1, 8), [40, 14, 16, 20, 35, 14, 18]):
        ws.column_dimensions[get_column_letter(c)].width = width

    for sheet_name, payload in result_sheets.items():
        ws = out_wb.create_sheet(safe_sheet_name(sheet_name))
        headers = payload["headers"]
        rows = payload["rows"]
        for c, h in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=c, value=h)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL_2B
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.freeze_panes = "A2"
        for r, row in enumerate(rows, start=2):
            for c, val in enumerate(row, start=1):
                ws.cell(row=r, column=c, value=val)
        ws.column_dimensions["A"].width = 14
        for c in range(2, len(headers) + 1):
            ws.column_dimensions[get_column_letter(c)].width = 18

    buf = io.BytesIO()
    out_wb.save(buf)
    buf.seek(0)
    return buf


def render_gstr2b():
    tool_hero(
        "\U0001F4CA", "GSTR-2B Consolidator",
        "Combine your GST purchase data in seconds",
        "Upload multiple GSTR-2B Excel files and automatically consolidate them "
        "into RCM / B2B / B2BA / B2B-CDNR / Pivot / Net / IMPG / IMPGSEZ / IMPGA "
        "sheets, with any other sheet types appended after. The 'Read me' sheet "
        "is never consolidated.",
    )

    with st.container(border=True):
        st.markdown("**\u2B06 Upload your GSTR-2B files**")
        uploaded_files = st.file_uploader(
            "Upload GSTR-2B excel files", type=["xlsx"], accept_multiple_files=True,
            key="gstr2b_uploader", label_visibility="collapsed",
        )
        st.caption("Supports .XLSX \u2022 Multiple files")

    if not uploaded_files:
        st.info("Upload your GSTR-2B excel files to get started.")
        return

    st.markdown(f"**Uploaded Files** \u2014 {len(uploaded_files)} file(s) selected")
    if len(uploaded_files) < 2:
        st.info("Tip: upload two or more files to actually see consolidation across periods.")

    with st.spinner("Reading and consolidating files..."):
        result_sheets, file_summaries, warnings = engine2b.process_files(uploaded_files)

    for w in warnings:
        st.warning(w)

    st.success("\u2705 Consolidation complete \u2014 your workbook is ready below.")

    st.subheader("Files processed")
    st.dataframe(
        [
            {
                "File": fs["filename"], "Month": fs["period"], "FY": fs["financial_year"],
                "GSTIN": fs["gstin"], "Legal Name": fs["legal_name"],
                "Sheets": fs["sheets"], "Rows added": fs["rows"],
            }
            for fs in file_summaries
        ],
        width='stretch', hide_index=True,
    )

    total_rows = sum(len(p["rows"]) for p in result_sheets.values())
    non_empty_sheets = sum(1 for p in result_sheets.values() if p["rows"])
    c1, c2, c3 = st.columns(3)
    c1.metric("Files Consolidated", len(file_summaries))
    c2.metric("Sheets With Data", f"{non_empty_sheets} / {len(result_sheets)}")
    c3.metric("Total Data Rows", total_rows)

    import pandas as pd

    if "Pivot" in result_sheets and result_sheets["Pivot"]["rows"]:
        pivot_headers = result_sheets["Pivot"]["headers"]
        grand_total_row = next((r for r in result_sheets["Pivot"]["rows"] if r[0] == "Grand Total"), None)
        if grand_total_row:
            igst, cgst, sgst, cess = grand_total_row[1], grand_total_row[2], grand_total_row[3], grand_total_row[4]
            t1, t2, t3, t4, t5 = st.columns(5)
            t1.metric("Total IGST (\u20B9)", f"{igst:,.2f}")
            t2.metric("Total CGST (\u20B9)", f"{cgst:,.2f}")
            t3.metric("Total SGST (\u20B9)", f"{sgst:,.2f}")
            t4.metric("Total Cess (\u20B9)", f"{cess:,.2f}")
            t5.metric("Total ITC (All Taxes) (\u20B9)", f"{igst + cgst + sgst + cess:,.2f}")

        st.subheader("Month-wise ITC totals (Pivot)")
        pdf = pd.DataFrame(result_sheets["Pivot"]["rows"], columns=pivot_headers)
        st.dataframe(pdf, width='stretch', hide_index=True)

    sheets_with_data = [(name, len(p["rows"])) for name, p in result_sheets.items() if p["rows"] and name != "Pivot"]
    if sheets_with_data:
        st.subheader("Sheet Summary")
        cols = st.columns(4)
        for i, (name, count) in enumerate(sheets_with_data):
            with cols[i % 4]:
                with st.container(border=True):
                    st.markdown(f"**{name}**")
                    st.caption(f"{count:,} row(s)")

    st.subheader("Preview a sheet")
    sheet_choice = st.selectbox(
        "Pick a sheet to preview",
        options=[s for s in result_sheets.keys() if result_sheets[s]["rows"]] or list(result_sheets.keys()),
        key="gstr2b_sheet_choice",
    )
    if sheet_choice:
        payload = result_sheets[sheet_choice]
        if payload["rows"]:
            df = pd.DataFrame(payload["rows"], columns=payload["headers"])
            st.dataframe(df, width='stretch', height=350)
        else:
            st.caption("No data rows in this sheet across the uploaded files.")

    output_buf = build_2b_workbook(result_sheets, file_summaries)
    st.download_button(
        label="\u2B07\uFE0F Download Consolidated Excel",
        data=output_buf,
        file_name=f"GSTR2B_Consolidated_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        key="gstr2b_download",
        type="primary",
        width='stretch',
    )


# ==========================================================================
# GSTR-3B tool  (logic calls unchanged from the previous version)
# ==========================================================================

OUTPUT_FILL_3B = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
INPUT_FILL_3B = PatternFill(start_color="6E2C00", end_color="6E2C00", fill_type="solid")
OTHER_FILL_3B = PatternFill(start_color="4A4A4A", end_color="4A4A4A", fill_type="solid")

TAB_COLORS_3B = {}
for _name in engine3b.OUTPUT_SHEETS:
    TAB_COLORS_3B[_name] = "1F4E78"
for _name in engine3b.INPUT_SHEETS:
    TAB_COLORS_3B[_name] = "6E2C00"
for _name in engine3b.OTHER_SHEETS:
    TAB_COLORS_3B[_name] = "4A4A4A"


def build_3b_workbook(result_sheets, file_summaries):
    out_wb = openpyxl.Workbook()
    out_wb.remove(out_wb.active)

    ws = out_wb.create_sheet("Filing Summary")
    ws["A1"] = "GSTR-3B Consolidation Summary"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Generated on {datetime.now().strftime('%d-%b-%Y %H:%M')}"
    ws["A2"].font = SUBTITLE_FONT

    headers = ["File Name", "Month", "Financial Year", "GSTIN", "Legal Name", "ARN", "Date of Filing", "Rows Extracted"]
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=c, value=h)
        cell.font = HEADER_FONT
        cell.fill = OUTPUT_FILL_3B
    for r, fs in enumerate(file_summaries, start=5):
        ws.cell(row=r, column=1, value=fs["filename"])
        ws.cell(row=r, column=2, value=fs["period"])
        ws.cell(row=r, column=3, value=fs["financial_year"])
        ws.cell(row=r, column=4, value=fs["gstin"])
        ws.cell(row=r, column=5, value=fs["legal_name"])
        ws.cell(row=r, column=6, value=fs["arn"])
        ws.cell(row=r, column=7, value=fs["date_of_filing"])
        ws.cell(row=r, column=8, value=fs["rows"])
    for c, width in zip(range(1, 9), [40, 12, 14, 20, 35, 20, 16, 14]):
        ws.column_dimensions[get_column_letter(c)].width = width

    for sheet_name, payload in result_sheets.items():
        ws = out_wb.create_sheet(safe_sheet_name(sheet_name))
        ws.sheet_properties.tabColor = TAB_COLORS_3B.get(sheet_name, "1F4E78")
        fill = OUTPUT_FILL_3B if sheet_name in engine3b.OUTPUT_SHEETS else (
            INPUT_FILL_3B if sheet_name in engine3b.INPUT_SHEETS else OTHER_FILL_3B)

        headers = payload["headers"]
        rows = payload["rows"]
        for c, h in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=c, value=h)
            cell.font = HEADER_FONT
            cell.fill = fill
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.freeze_panes = "A2"
        for r, row in enumerate(rows, start=2):
            for c, val in enumerate(row, start=1):
                ws.cell(row=r, column=c, value=val)
        ws.column_dimensions["A"].width = 12
        ws.column_dimensions["B"].width = 42
        for c in range(3, len(headers) + 1):
            ws.column_dimensions[get_column_letter(c)].width = 18

    buf = io.BytesIO()
    out_wb.save(buf)
    buf.seek(0)
    return buf


def render_gstr3b():
    tool_hero(
        "\U0001F4CB", "GSTR-3B Consolidator",
        "Consolidate your return data, bifurcated by Output and Input",
        "Upload multiple GSTR-3B PDFs (as downloaded from the GST portal). Every "
        "table becomes its own annexure sheet — bifurcated into Output (outward "
        "supplies, export/SEZ, RCM liability) and Input (ITC) — with each row "
        "tagged by Month and consolidated across all uploaded periods.",
    )

    with st.container(border=True):
        st.markdown("**\u2B06 Upload your GSTR-3B filings**")
        uploaded_files = st.file_uploader(
            "Upload GSTR-3B PDF files", type=["pdf"], accept_multiple_files=True,
            key="gstr3b_uploader", label_visibility="collapsed",
        )
        st.caption("Supports .PDF \u2022 Multiple files")

    if not uploaded_files:
        st.info("Upload your GSTR-3B PDF files to get started.")
        return

    st.markdown(f"**Uploaded Files** \u2014 {len(uploaded_files)} file(s) selected")
    if len(uploaded_files) < 2:
        st.info("Tip: upload two or more months to actually see consolidation across periods.")

    with st.spinner("Reading and parsing PDFs..."):
        result_sheets, file_summaries, warnings = engine3b.process_files(uploaded_files)

    for w in warnings:
        st.warning(w)

    st.success("\u2705 Consolidation complete \u2014 your workbook is ready below.")

    st.subheader("Files processed")
    st.dataframe(
        [
            {
                "File": fs["filename"], "Month": fs["period"], "FY": fs["financial_year"],
                "GSTIN": fs["gstin"], "Legal Name": fs["legal_name"],
                "ARN": fs["arn"], "Filed on": fs["date_of_filing"], "Rows extracted": fs["rows"],
            }
            for fs in file_summaries
        ],
        width='stretch', hide_index=True,
    )

    total_rows = sum(len(p["rows"]) for p in result_sheets.values())
    c1, c2, c3 = st.columns(3)
    c1.metric("Filings Processed", len(file_summaries))
    c2.metric("Annexure Sheets", len(result_sheets))
    c3.metric("Total Rows Extracted", total_rows)

    import pandas as pd

    tab_output, tab_input, tab_other = st.tabs(["\U0001F4E4 Output", "\U0001F4E5 Input (ITC)", "\U0001F4C4 Other"])

    def render_group(container, sheet_names):
        with container:
            for name in sheet_names:
                payload = result_sheets[name]
                if payload["rows"]:
                    st.markdown(f"**{name}**")
                    df = pd.DataFrame(payload["rows"], columns=payload["headers"])
                    st.dataframe(df, width='stretch', hide_index=True)

    render_group(tab_output, engine3b.OUTPUT_SHEETS)
    render_group(tab_input, engine3b.INPUT_SHEETS)
    render_group(tab_other, engine3b.OTHER_SHEETS)

    output_buf = build_3b_workbook(result_sheets, file_summaries)
    st.download_button(
        label="\u2B07\uFE0F Download Consolidated Excel",
        data=output_buf,
        file_name=f"GSTR3B_Consolidated_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        key="gstr3b_download",
        type="primary",
        width='stretch',
    )


# ==========================================================================
# PDF -> Excel tool  (logic calls unchanged from the previous version)
# ==========================================================================

def render_pdf2excel():
    tool_hero(
        "\U0001F4C4", "PDF to Excel Converter",
        "Extract tables from any PDF in a few clicks",
        "Upload one or more PDFs. Every table pdfplumber can detect is extracted "
        "onto its own sheet; multi-page tables (like a bank statement or ledger) "
        "are automatically merged into one sheet, and pages with no table fall "
        "back to a plain-text sheet so nothing gets silently dropped.",
    )

    with st.container(border=True):
        st.markdown("**\u2B06 Upload your PDF files**")
        uploaded_files = st.file_uploader(
            "Upload PDF files", type=["pdf"], accept_multiple_files=True,
            key="pdf2excel_uploader", label_visibility="collapsed",
        )
        st.caption("Supports .PDF \u2022 Multiple files")

    if not uploaded_files:
        st.info("Upload one or more PDF files to get started.")
        return

    st.markdown(f"**Uploaded Files** \u2014 {len(uploaded_files)} file(s) selected")

    col1, col2 = st.columns(2)
    with col1:
        merge_across = st.checkbox(
            "Merge multi-page tables into one sheet",
            value=True, key="pdf2excel_merge_across",
            help="On (default): consecutive tables with the same column count (e.g. a bank statement or ledger spanning many pages) are combined into a single sheet, with any repeated header row stripped. Off: every table gets its own per-page sheet.",
        )
    with col2:
        merge_per_page = st.checkbox(
            "Merge multiple tables on the same page",
            value=False, key="pdf2excel_merge_per_page",
            help="On: if a single page has more than one detected table, stack them into one sheet separated by a blank row. Off (default): each gets its own sheet.",
        )

    with st.spinner("Extracting tables..."):
        all_file_sheets = []
        total_tables = 0
        for uf in uploaded_files:
            sheets = pdf2excel.extract_pdf_to_sheets(
                uf.getvalue(), merge_tables_per_page=merge_per_page, merge_across_pages=merge_across,
            )
            all_file_sheets.append((uf.name, sheets))
            total_tables += len(sheets)

    st.success("\u2705 Extraction complete \u2014 your workbook is ready below.")

    c1, c2 = st.columns(2)
    c1.metric("Files Processed", len(uploaded_files))
    c2.metric("Sheets Extracted", total_tables)

    st.subheader("Preview")
    flat_options = []
    for filename, sheets in all_file_sheets:
        for sheet_name, kind, rows in sheets:
            flat_options.append((f"{filename} — {sheet_name}", kind, rows))
    if flat_options:
        labels = [o[0] for o in flat_options]
        choice = st.selectbox("Pick a sheet to preview", labels, key="pdf2excel_preview_choice")
        _, kind, rows = flat_options[labels.index(choice)]
        if kind == "table" and rows:
            import pandas as pd
            width = max(len(r) for r in rows)
            padded = [list(r) + [None] * (width - len(r)) for r in rows]
            df = pd.DataFrame(padded[1:], columns=[str(c) for c in padded[0]]) if len(padded) > 1 else pd.DataFrame(padded)
            st.dataframe(df, width='stretch', height=300)
        else:
            st.text("\n".join(rows[:50]))
    else:
        st.caption("No tables or text could be extracted from the uploaded PDF(s).")

    output_buf = pdf2excel.build_workbook(all_file_sheets)
    st.download_button(
        label="\u2B07\uFE0F Download Excel Workbook",
        data=output_buf,
        file_name=f"PDF_to_Excel_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        key="pdf2excel_download",
        type="primary",
        width='stretch',
    )


# ==========================================================================
# Excel -> PDF tool  (logic calls unchanged from the previous version)
# ==========================================================================

def render_excel2pdf():
    tool_hero(
        "\U0001F4D1", "Excel to PDF Converter",
        "Turn any workbook into a clean, professional PDF",
        "Upload one or more Excel workbooks. Every sheet is rendered as a table "
        "in a landscape PDF, one section per sheet, auto-paginated for long sheets.",
    )

    with st.container(border=True):
        st.markdown("**\u2B06 Upload your Excel files**")
        uploaded_files = st.file_uploader(
            "Upload Excel files", type=["xlsx"], accept_multiple_files=True,
            key="excel2pdf_uploader", label_visibility="collapsed",
        )
        st.caption("Supports .XLSX \u2022 Multiple files")

    if not uploaded_files:
        st.info("Upload one or more .xlsx files to get started.")
        return

    st.markdown(f"**Uploaded Files** \u2014 {len(uploaded_files)} file(s) selected")

    for uf in uploaded_files:
        with st.spinner(f"Converting {uf.name}..."):
            try:
                pdf_buf = excel2pdf.workbook_to_pdf(uf.getvalue(), uf.name)
            except Exception as e:
                st.error(f"Could not convert '{uf.name}': {e}")
                continue
        st.success(f"\u2705 {uf.name} converted \u2014 ready to download.")
        st.download_button(
            label=f"\u2B07\uFE0F Download {uf.name.rsplit('.', 1)[0]}.pdf",
            data=pdf_buf,
            file_name=f"{uf.name.rsplit('.', 1)[0]}.pdf",
            mime="application/pdf",
            key=f"excel2pdf_download_{uf.name}",
            type="primary",
        )


# ==========================================================================
# JSON -> Excel tool  (logic calls unchanged from the previous version)
# ==========================================================================

def render_json2excel():
    tool_hero(
        "\U0001F504", "JSON to Excel Converter",
        "Convert GST JSON exports into structured spreadsheets",
        "Upload one or more JSON files. Top-level keys become sheets; nested "
        "objects flatten into dotted column names; nested arrays of records "
        "(e.g. GST-style invoice -> item structures) are exploded into rows "
        "with their parent fields repeated.",
    )

    with st.container(border=True):
        st.markdown("**\u2B06 Upload your JSON files**")
        uploaded_files = st.file_uploader(
            "Upload JSON files", type=["json"], accept_multiple_files=True,
            key="json2excel_uploader", label_visibility="collapsed",
        )
        st.caption("Supports .JSON \u2022 Multiple files")

    if not uploaded_files:
        st.info("Upload one or more .json files to get started.")
        return

    st.markdown(f"**Uploaded Files** \u2014 {len(uploaded_files)} file(s) selected")

    with st.spinner("Parsing JSON..."):
        all_file_sheets = []
        warnings = []
        for uf in uploaded_files:
            try:
                sheets = json2excel.parse_json_to_sheets(uf.getvalue())
                all_file_sheets.append((uf.name, sheets))
            except Exception as e:
                warnings.append(f"Could not parse '{uf.name}': {e}")

    for w in warnings:
        st.warning(w)

    if not all_file_sheets:
        return

    st.success("\u2705 Conversion complete \u2014 your workbook is ready below.")

    total_sheets = sum(len(s) for _, s in all_file_sheets)
    c1, c2 = st.columns(2)
    c1.metric("Files Processed", len(all_file_sheets))
    c2.metric("Sheets Produced", total_sheets)

    st.subheader("Preview")
    flat_options = []
    for filename, sheets in all_file_sheets:
        for sheet_name, df in sheets:
            flat_options.append((f"{filename} — {sheet_name}", df))
    if flat_options:
        labels = [o[0] for o in flat_options]
        choice = st.selectbox("Pick a sheet to preview", labels, key="json2excel_preview_choice")
        df = flat_options[labels.index(choice)][1]
        st.dataframe(df, width='stretch', height=300)

    output_buf = json2excel.build_workbook(all_file_sheets)
    st.download_button(
        label="\u2B07\uFE0F Download Excel Workbook",
        data=output_buf,
        file_name=f"JSON_to_Excel_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        key="json2excel_download",
        type="primary",
        width='stretch',
    )


# ==========================================================================
# Category / tool registry — the ONLY place that needs an entry when a new
# tool is added later. "implemented": False tools show as "Coming Soon"
# and are never clickable, so nothing fake ever appears functional.
# ==========================================================================

CATEGORIES = [
    {
        "key": "gst_compliance", "icon": "\U0001F9FE", "name": "GST Compliance",
        "desc": "Manage and process GST return data",
        "page_subtitle": "GST return processing and compliance tools",
        "tools": [
            {"key": "gst_compliance::gstr1", "icon": "\U0001F4CA", "name": "GSTR-1", "desc": "Outward supplies", "implemented": False},
            {"key": "gst_compliance::gstr2a", "icon": "\U0001F4E5", "name": "GSTR-2A", "desc": "Purchase data", "implemented": False},
            {"key": "gst_compliance::gstr2b", "icon": "\U0001F4CA", "name": "GSTR-2B", "desc": "ITC statement", "implemented": True, "render": render_gstr2b},
            {"key": "gst_compliance::gstr3b", "icon": "\U0001F4CB", "name": "GSTR-3B", "desc": "Summary return", "implemented": True, "render": render_gstr3b},
        ],
    },
    {
        "key": "qrmp", "icon": "\U0001F4CB", "name": "QRMP",
        "desc": "Tools for Quarterly Return Monthly Payment taxpayers",
        "page_subtitle": "Quarterly Return Monthly Payment",
        "tools": [
            {"key": "qrmp::gstr1", "icon": "\U0001F4CA", "name": "GSTR-1", "desc": "Outward supplies (QRMP)", "implemented": False},
            {"key": "qrmp::gstr2a", "icon": "\U0001F4E5", "name": "GSTR-2A", "desc": "Purchase data (QRMP)", "implemented": False},
            {"key": "qrmp::gstr2b", "icon": "\U0001F4CA", "name": "GSTR-2B", "desc": "ITC statement (QRMP)", "implemented": False},
            {"key": "qrmp::gstr3b", "icon": "\U0001F4CB", "name": "GSTR-3B", "desc": "Summary return (QRMP)", "implemented": False},
        ],
    },
    {
        "key": "reconciliation", "icon": "\U0001F504", "name": "Reconciliation",
        "desc": "Compare GST data and identify differences",
        "page_subtitle": "Reconcile books and returns to identify mismatches",
        "tools": [
            {"key": "recon::books_2a", "icon": "\U0001F504", "name": "Books vs GSTR-2A", "desc": "Reconcile purchase register with 2A", "implemented": False},
            {"key": "recon::books_2b", "icon": "\U0001F504", "name": "Books vs GSTR-2B", "desc": "Reconcile purchase register with 2B", "implemented": False},
            {"key": "recon::2a_2b", "icon": "\U0001F504", "name": "GSTR-2A vs GSTR-2B", "desc": "Compare 2A and 2B", "implemented": False},
            {"key": "recon::3b_books", "icon": "\U0001F504", "name": "GSTR-3B vs Books", "desc": "Reconcile summary return with books", "implemented": False},
        ],
    },
    {
        "key": "file_conversion", "icon": "\U0001F4C4", "name": "File Conversion",
        "desc": "Convert and structure GST-related files",
        "page_subtitle": "Utilities to convert between PDF, Excel and JSON",
        "tools": [
            {"key": "conv::pdf2excel", "icon": "\U0001F4C4", "name": "PDF \u2192 Excel", "desc": "Extract tables from PDFs", "implemented": True, "render": render_pdf2excel},
            {"key": "conv::excel2pdf", "icon": "\U0001F4D1", "name": "Excel \u2192 PDF", "desc": "Create professional PDFs", "implemented": True, "render": render_excel2pdf},
            {"key": "conv::json2excel", "icon": "\U0001F504", "name": "JSON \u2192 Excel", "desc": "Convert GST JSON files", "implemented": True, "render": render_json2excel},
        ],
    },
    {
        "key": "audit_tools", "icon": "\U0001F50D", "name": "Audit Tools",
        "desc": "Tools designed for audit and financial review",
        "page_subtitle": "Future tools will be added here",
        "tools": [],
    },
]

CATEGORY_BY_KEY = {c["key"]: c for c in CATEGORIES}
TOOL_RENDER_MAP = {t["key"]: t["render"] for c in CATEGORIES for t in c["tools"] if t.get("implemented")}


# ==========================================================================
# Dashboard + category landing pages
# ==========================================================================

def render_dashboard():
    st.markdown(
        '<div style="padding: 1rem 0 0.25rem 0;">'
        '<div style="font-size:2rem; font-weight:800;">Welcome to Thaanvi\'s Audit Suite \U0001F44B</div>'
        f'<div style="font-size:1.1rem; color:{TEAL}; font-weight:600; margin-top:6px;">GST Compliance &amp; Audit Tools</div>'
        '<div style="font-size:0.95rem; color:#5B6B7C; margin-top:2px;">Simplify. Reconcile. Review. Report.</div>'
        '</div>',
        unsafe_allow_html=True,
    )
    st.write("")

    cols = st.columns(3)
    for i, cat in enumerate(CATEGORIES):
        with cols[i % 3]:
            with st.container(border=True):
                st.markdown(f"### {cat['icon']} {cat['name']}")
                st.caption(cat["desc"])
                if st.button("Explore \u2192", key=f"open_cat_{cat['key']}", width='stretch'):
                    st.session_state.nav = {"level": "category", "category": cat["key"]}
                    st.rerun()
        if i % 3 == 2:
            st.write("")


def render_category(category_key):
    cat = CATEGORY_BY_KEY[category_key]
    st.markdown(f'<div class="gcs-hero-title">{cat["icon"]} {cat["name"]}</div>', unsafe_allow_html=True)
    st.markdown(f'<div class="gcs-hero-desc">{cat["page_subtitle"]}</div>', unsafe_allow_html=True)

    if not cat["tools"]:
        st.info("Future tools will be added here.")
        return

    cols = st.columns(3)
    for i, tool in enumerate(cat["tools"]):
        with cols[i % 3]:
            with st.container(border=True):
                if tool["implemented"]:
                    st.markdown(f"**{tool['icon']} {tool['name']}**")
                else:
                    st.markdown(f"**{tool['icon']} {tool['name']}** <span class='gcs-soon-badge'>COMING SOON</span>", unsafe_allow_html=True)
                st.caption(tool["desc"])
                if tool["implemented"]:
                    if st.button("Open Tool", key=f"opentool_{tool['key']}", width='stretch'):
                        st.session_state.nav = {"level": "tool", "category": category_key, "tool": tool["key"]}
                        st.rerun()
                else:
                    st.button("Coming Soon", key=f"soontool_{tool['key']}", width='stretch', disabled=True)
        if i % 3 == 2:
            st.write("")


# ==========================================================================
# Sidebar navigation
# ==========================================================================

if "nav" not in st.session_state:
    st.session_state.nav = {"level": "dashboard"}

nav = st.session_state.nav

st.sidebar.markdown(
    '<div style="font-size:1.1rem; font-weight:800; padding: 4px 0 0 0; line-height:1.3;">THAANVI\'S<br>AUDIT SUITE</div>',
    unsafe_allow_html=True,
)
st.sidebar.markdown('<div style="font-size:0.75rem; color:#A9BCCE; margin-bottom:10px;">GST Compliance &amp; Audit Tools</div>', unsafe_allow_html=True)

if st.sidebar.button("\U0001F3E0 Dashboard", key="nav_dashboard", width='stretch',
                      type="primary" if nav["level"] == "dashboard" else "secondary"):
    st.session_state.nav = {"level": "dashboard"}
    st.rerun()

st.sidebar.markdown("---")

for cat in CATEGORIES:
    is_active_cat = nav.get("category") == cat["key"] and nav["level"] in ("category", "tool")
    if st.sidebar.button(
        f"{cat['icon']} {cat['name']}", key=f"nav_cat_{cat['key']}", width='stretch',
        type="primary" if (is_active_cat and nav["level"] == "category") else "secondary",
    ):
        st.session_state.nav = {"level": "category", "category": cat["key"]}
        st.rerun()

    if is_active_cat:
        if not cat["tools"]:
            st.sidebar.caption("\u3000Future tools will be added here.")
        for tool in cat["tools"]:
            if tool["implemented"]:
                is_active_tool = nav["level"] == "tool" and nav.get("tool") == tool["key"]
                if st.sidebar.button(
                    f"\u3000{tool['icon']} {tool['name']}", key=f"nav_tool_{tool['key']}", width='stretch',
                    type="primary" if is_active_tool else "secondary",
                ):
                    st.session_state.nav = {"level": "tool", "category": cat["key"], "tool": tool["key"]}
                    st.rerun()
            else:
                st.sidebar.button(
                    f"\u3000{tool['icon']} {tool['name']} \u00b7 Soon", key=f"nav_tool_{tool['key']}",
                    width='stretch', disabled=True,
                )

st.sidebar.markdown("---")
st.sidebar.markdown("### Appearance")
st.sidebar.radio("Theme", ["Light", "Dark"], horizontal=True, key="theme_mode_radio")


# ==========================================================================
# Page dispatch
# ==========================================================================

if nav["level"] == "dashboard":
    render_dashboard()
elif nav["level"] == "category":
    top_bar()
    render_category(nav["category"])
elif nav["level"] == "tool":
    top_bar()
    cat = CATEGORY_BY_KEY[nav["category"]]
    if st.button(f"\u2190 Back to {cat['name']}", key="tool_back_button"):
        st.session_state.nav = {"level": "category", "category": nav["category"]}
        st.rerun()
    TOOL_RENDER_MAP[nav["tool"]]()
