"""
GSTR-2B Consolidator
---------------------
Upload multiple GSTR-2B excel exports (as downloaded from the GST portal) and
get back ONE workbook consolidated in the standard working-papers format:

  - RCM        : reverse-charge invoices pulled out of B2B + B2BA
  - B2B        : regular B2B invoices (reverse-charge rows moved to RCM)
  - B2BA       : amended B2B invoices (reverse-charge rows moved to RCM)
  - B2B-CDNR   : credit/debit notes
  - Pivot      : month-wise IGST/CGST/SGST/Cess totals (from Net)
  - Net        : one consolidated ITC register (B2B+CDNR+IMPG+IMPGSEZ+B2BA+IMPGA)
  - IMPG / IMPGSEZ / IMPGA : import entries

Every other GSTR-2B sheet type present in the uploaded files (ITC Available,
ISD, ECO, rejected/amendment variants, etc.) is also consolidated and
appended after these, so nothing from the source files is lost.

"Read me" sheet is never consolidated.
"""

import io
from datetime import datetime

import streamlit as st
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

import engine

st.set_page_config(page_title="GSTR-2B Consolidator", page_icon="\U0001F4CA", layout="wide")

# --------------------------------------------------------------------------
# Workbook builder
# --------------------------------------------------------------------------

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(bold=True, size=14, color="1F4E78")
SUBTITLE_FONT = Font(italic=True, size=10, color="666666")


def build_output_workbook(result_sheets, file_summaries):
    out_wb = openpyxl.Workbook()
    out_wb.remove(out_wb.active)

    # ---- Summary sheet first ----
    ws = out_wb.create_sheet("Consolidation Summary")
    ws["A1"] = "GSTR-2B Consolidation Summary"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Generated on {datetime.now().strftime('%d-%b-%Y %H:%M')}"
    ws["A2"].font = SUBTITLE_FONT

    headers = ["File Name", "Month", "Financial Year", "GSTIN", "Legal Name", "Sheets Found", "Rows Consolidated"]
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=c, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    for r, fs in enumerate(file_summaries, start=5):
        ws.cell(row=r, column=1, value=fs["filename"])
        ws.cell(row=r, column=2, value=fs["period"])
        ws.cell(row=r, column=3, value=fs["financial_year"])
        ws.cell(row=r, column=4, value=fs["gstin"])
        ws.cell(row=r, column=5, value=fs["legal_name"])
        ws.cell(row=r, column=6, value=fs["sheets"])
        ws.cell(row=r, column=7, value=fs["rows"])
    for c, width in zip(range(1, 8), [40, 14, 16, 20, 35, 14, 18]):
        ws.column_dimensions[get_column_letter(c)].width = width

    # ---- One sheet per consolidated type, in the curated order ----
    for sheet_name, payload in result_sheets.items():
        safe_name = sheet_name[:31]
        ws = out_wb.create_sheet(safe_name)
        headers = payload["headers"]
        rows = payload["rows"]

        for c, h in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=c, value=h)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.freeze_panes = "A2"

        for r, row in enumerate(rows, start=2):
            for c, val in enumerate(row, start=1):
                ws.cell(row=r, column=c, value=val)

        ws.column_dimensions["A"].width = 14
        for c in range(2, len(headers) + 1):
            ws.column_dimensions[get_column_letter(c)].width = 18

    buf = io.BytesIO()
    out_wb.save(buf)
    buf.seek(0)
    return buf


# --------------------------------------------------------------------------
# Streamlit UI
# --------------------------------------------------------------------------

st.title("\U0001F4CA GSTR-2B Consolidator")
st.caption(
    "Upload two or more GSTR-2B excel files (as exported from the GST portal). "
    "They'll be consolidated into RCM / B2B / B2BA / B2B-CDNR / Pivot / Net / "
    "IMPG / IMPGSEZ / IMPGA sheets, with any other sheet types appended after. "
    "The 'Read me' sheet is never consolidated."
)

uploaded_files = st.file_uploader(
    "Upload GSTR-2B excel files",
    type=["xlsx"],
    accept_multiple_files=True,
)

if uploaded_files:
    if len(uploaded_files) < 2:
        st.info("Tip: upload two or more files to actually see consolidation across periods.")

    with st.spinner("Reading and consolidating files..."):
        result_sheets, file_summaries, warnings = engine.process_files(uploaded_files)

    if warnings:
        for w in warnings:
            st.warning(w)

    st.subheader("Files processed")
    st.dataframe(
        [
            {
                "File": fs["filename"],
                "Month": fs["period"],
                "FY": fs["financial_year"],
                "GSTIN": fs["gstin"],
                "Legal Name": fs["legal_name"],
                "Sheets": fs["sheets"],
                "Rows added": fs["rows"],
            }
            for fs in file_summaries
        ],
        use_container_width=True,
        hide_index=True,
    )

    total_rows = sum(len(p["rows"]) for p in result_sheets.values())
    non_empty_sheets = sum(1 for p in result_sheets.values() if p["rows"])
    c1, c2, c3 = st.columns(3)
    c1.metric("Files consolidated", len(file_summaries))
    c2.metric("Sheets with data", f"{non_empty_sheets} / {len(result_sheets)}")
    c3.metric("Total data rows", total_rows)

    if "Pivot" in result_sheets and result_sheets["Pivot"]["rows"]:
        st.subheader("Month-wise ITC totals (Pivot)")
        import pandas as pd
        pdf = pd.DataFrame(result_sheets["Pivot"]["rows"], columns=result_sheets["Pivot"]["headers"])
        st.dataframe(pdf, use_container_width=True, hide_index=True)

    st.subheader("Preview a sheet")
    sheet_choice = st.selectbox(
        "Pick a sheet to preview",
        options=[s for s in result_sheets.keys() if result_sheets[s]["rows"]] or list(result_sheets.keys()),
    )
    if sheet_choice:
        payload = result_sheets[sheet_choice]
        if payload["rows"]:
            import pandas as pd
            df = pd.DataFrame(payload["rows"], columns=payload["headers"])
            st.dataframe(df, use_container_width=True, height=350)
        else:
            st.caption("No data rows in this sheet across the uploaded files.")

    output_buf = build_output_workbook(result_sheets, file_summaries)
    st.download_button(
        label="\u2B07\uFE0F Download consolidated workbook",
        data=output_buf,
        file_name=f"GSTR2B_Consolidated_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
else:
    st.info("Upload your GSTR-2B excel files to get started.")
