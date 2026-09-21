"""
JSON -> Excel Converter
--------------------------
Flattens an uploaded JSON file into one or more sheets.

Nested structures are fully flattened, not just one level deep — this
matters for things like GST JSON exports where a record can be nested
several arrays deep (e.g. b2b -> inv -> items). Nested dicts become
dot-notation columns (e.g. "address.city"); nested arrays-of-objects are
exploded into repeated rows (so parent fields repeat once per child
record, the usual "flatten to tabular" behaviour); arrays of plain values
(strings/numbers) are joined into a single comma-separated cell rather
than exploded, to avoid uncontrolled row blow-up for simple tag lists.

  - top-level dict: each key becomes its own sheet
  - top-level list: a single "Data" sheet
"""

import io
import json

import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
INVALID_SHEET_CHARS = str.maketrans({c: "-" for c in r'\/?*[]:'})


def safe_sheet_name(name):
    return name.translate(INVALID_SHEET_CHARS)[:31] or "Sheet"


def _flatten_record(record, parent_key=""):
    """
    Fully flatten one dict into a list of flat dicts (>1 if it contains a
    nested list-of-objects, since those get exploded into separate rows).
    """
    if not isinstance(record, dict):
        key = parent_key or "value"
        return [{key: record}]

    scalars = {}
    exploded_groups = []  # list of list_of_flat_dicts

    for k, v in record.items():
        full_key = f"{parent_key}.{k}" if parent_key else str(k)
        if isinstance(v, dict):
            sub_rows = _flatten_record(v, full_key)
            exploded_groups.append(sub_rows)
        elif isinstance(v, list):
            if len(v) and all(isinstance(x, dict) for x in v):
                sub_rows = []
                for item in v:
                    sub_rows.extend(_flatten_record(item, full_key))
                exploded_groups.append(sub_rows)
            elif v:
                scalars[full_key] = ", ".join(str(x) for x in v)
            else:
                scalars[full_key] = None
        else:
            scalars[full_key] = v

    rows = [dict(scalars)]
    for sub_rows in exploded_groups:
        if not sub_rows:
            continue
        new_rows = []
        for base in rows:
            for sub in sub_rows:
                merged = dict(base)
                merged.update(sub)
                new_rows.append(merged)
        rows = new_rows if new_rows else rows

    return rows


def _value_to_df(value):
    """Turn an arbitrary JSON value into a DataFrame as sensibly as possible."""
    if isinstance(value, list):
        if not value:
            return pd.DataFrame()
        if all(isinstance(v, dict) for v in value):
            all_rows = []
            for rec in value:
                all_rows.extend(_flatten_record(rec))
            return pd.DataFrame(all_rows)
        return pd.DataFrame({"value": value})
    if isinstance(value, dict):
        return pd.DataFrame(_flatten_record(value))
    return pd.DataFrame({"value": [value]})


def parse_json_to_sheets(file_bytes):
    """Returns a list of (sheet_name, DataFrame) tuples."""
    data = json.loads(file_bytes.decode("utf-8"))
    sheets = []

    if isinstance(data, list):
        sheets.append(("Data", _value_to_df(data)))
    elif isinstance(data, dict):
        for key, value in data.items():
            df = _value_to_df(value)
            if not df.empty:
                sheets.append((str(key), df))
        if not sheets:
            sheets.append(("Data", pd.DataFrame(list(data.items()), columns=["Key", "Value"])))
    else:
        sheets.append(("Data", pd.DataFrame({"value": [data]})))

    return sheets


def build_workbook(all_file_sheets):
    """
    all_file_sheets: list of (filename, [(sheet_name, DataFrame), ...])
    """
    out_wb = openpyxl.Workbook()
    out_wb.remove(out_wb.active)
    multi_file = len(all_file_sheets) > 1

    used_names = set()

    def unique_name(base):
        name = safe_sheet_name(base)
        i = 2
        while name in used_names:
            suffix = f"_{i}"
            name = safe_sheet_name(base[:31 - len(suffix)] + suffix)
            i += 1
        used_names.add(name)
        return name

    for filename, sheets in all_file_sheets:
        prefix = filename.rsplit(".", 1)[0]
        for sheet_name, df in sheets:
            full_name = f"{prefix}_{sheet_name}" if multi_file else sheet_name
            ws = out_wb.create_sheet(unique_name(full_name))
            for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=True), start=1):
                for c_idx, val in enumerate(row, start=1):
                    if isinstance(val, (list, dict)):
                        val = json.dumps(val)
                    ws.cell(row=r_idx, column=c_idx, value=val)
            if len(df.columns):
                for c in range(1, len(df.columns) + 1):
                    cell = ws.cell(row=1, column=c)
                    cell.font = HEADER_FONT
                    cell.fill = HEADER_FILL
                ws.freeze_panes = "A2"

    if not out_wb.sheetnames:
        out_wb.create_sheet("Empty")

    buf = io.BytesIO()
    out_wb.save(buf)
    buf.seek(0)
    return buf
